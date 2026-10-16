import re
import sys
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

GREEN = RGBColor(0x16, 0x7A, 0x3A)
RED = RGBColor(0xC0, 0x1C, 0x1C)
GREY = RGBColor(0x6B, 0x72, 0x80)
HEADER_FILL = "DCE6F2"

# Friendly names for CSV/test-data columns; others are humanized automatically
DATA_LABELS = {
    "related_party": "MVNO (related party)",
    "code": "Product code",
    "category_name": "Category",
    "channel_name": "Channel",
    "correlation_group": "Correlation group",
    "skus": "SKUs",
}

# Keywords whose inner steps are request plumbing, not checks
DO_NOT_DESCEND = {"Get Product Offering", "Compatibility Check", "Get Cached Access Token"}


def _check_descriptions(data, args):
    """Plain-English description for each known check keyword."""
    channel = data.get("channel_name", "")
    return {
        "Validate Response Status": f"API returns HTTP status {args.get(1, '200')}",
        "Validate Response Not Empty": "Response body is not empty",
        "Validate Response Headers": (
            f"Response header '{args.get(1, '')}' contains '{args.get(2, '')}'"
        ),
        "Validate Response Time": f"Response time is within {args.get(1, '')} ms",
        "Validate Requested Product Offering Channels": (
            f"Channel filtering: the bundle, every product in it and their prices "
            f"are available on channel '{channel}'"
        ),
        "Validate Product Offering Schema": (
            "Every offering (bundle and products) has an id, code and name, "
            "and a valid lifecycle status"
        ),
        "Validate Product Offering Prices": (
            "Every price is a valid, non-negative number"
            + (f" and is available on channel '{channel}'" if channel else "")
        ),
        "Validate Product Offering Validity": (
            "Every offering has valid start/end dates and has not expired"
        ),
        "Validate Product Offering Response": (
            f"Returned offering matches the requested MVNO '{data.get('related_party', '')}', "
            f"product code '{data.get('code', '')}' and category '{data.get('category_name', '')}'"
        ),
        "Validate Compatibility Check Response": (
            f"Response contains the requested correlation group "
            f"'{data.get('correlation_group', '')}' and SKUs {data.get('skus', '')}"
        ),
        "Get Compatibility Result": "Compatibility result (true/false) is returned",
        "Get Access Token": "IAM server issues an access token (HTTP 200 with access_token)",
        "Should Not Be Empty": f"{_humanize(args.get(0, 'Value'))} is not empty",
        "Fail": "Final result check",
    }


def _humanize(name):
    name = re.sub(r"^\$\{(.*)\}$", r"\1", str(name))
    return name.replace("_", " ").strip().capitalize()


def format_duration(value):
    try:
        seconds = float(value or 0)
    except (TypeError, ValueError):
        return "-"
    if seconds < 60:
        return f"{seconds:.1f} s"
    minutes, remainder = divmod(int(round(seconds)), 60)
    return f"{minutes} m {remainder:02d} s"


def _shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_text(cell, text, color=None, bold=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def _fit_to_window(table):
    """Word's "AutoFit to Window": full page width, columns sized by content."""
    table.autofit = True
    properties = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        existing = properties.find(qn(tag))
        if existing is not None:
            properties.remove(existing)
    width = OxmlElement("w:tblW")
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    properties.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    properties.append(layout)


def _clear_cell_widths(document):
    """Let Word size columns by content: drop the fixed cell widths python-docx adds."""
    for table in document.tables:
        for cell_width in list(table._tbl.iter(qn("w:tcW"))):
            cell_width.set(qn("w:type"), "auto")
            cell_width.set(qn("w:w"), "0")


def _table(document, headers):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fit_to_window(table)
    for cell, header in zip(table.rows[0].cells, headers):
        _set_text(cell, header, bold=True)
        _shade(cell, HEADER_FILL)
    return table


def _add_row(table, values):
    cells = table.add_row().cells
    for cell, value in zip(cells, values):
        if isinstance(value, tuple):
            _set_text(cell, value[0], color=value[1], bold=True)
        else:
            _set_text(cell, value)
    return cells


def _status_label(status):
    return {
        "PASS": ("Passed", GREEN),
        "FAIL": ("Failed", RED),
        "NOT RUN": ("Not run", GREY),
        "SKIP": ("Skipped", GREY),
    }.get(status, (status.title(), GREY))


def _status(node):
    status = node.find("./status")
    return status.get("status", "UNKNOWN") if status is not None else "UNKNOWN"


def _resolve(value, data):
    """Turn a recorded argument like 'expected_channel=${channel_name}' into its value."""
    value = str(value or "")
    if re.match(r"^[A-Za-z_]+=", value):
        value = value.split("=", 1)[1]
    return re.sub(r"\$\{([^}]+)\}", lambda m: data.get(m.group(1), m.group(0)), value)


def _test_data(test):
    """Input values of the data-driven template keyword (e.g. channel_name=WEB)."""
    first = test.find("./kw")
    data = {}
    if first is not None:
        for arg in first.findall("./arg"):
            text = arg.text or ""
            match = re.match(r"^([A-Za-z_][A-Za-z0-9_]*)=(.*)$", text, re.S)
            if match:
                data[match.group(1)] = match.group(2)
    return data


def _failure_message(node):
    for message in node.findall("./msg"):
        if message.get("level") == "FAIL" and message.text:
            return message.text.strip()
    status = node.find("./status")
    return (status.text or "").strip() if status is not None else ""


def _pass_detail(node):
    """Summary a validator returned, e.g. 'Validated 1 group offering(s)...'."""
    for message in node.findall("./msg"):
        text = message.text or ""
        if message.get("level") == "INFO" and text.startswith("${") and " = " in text:
            value = text.split(" = ", 1)[1].strip()
            if value.startswith("Validated"):
                return value
    return ""


def _collect_checks(node, data, attempt=None, checks=None):
    checks = [] if checks is None else checks
    for child in node:
        if child.tag == "iter":
            iteration = next((v.text for v in child.findall("./var")), None)
            label = f"Attempt {int(iteration) + 1}" if iteration and iteration.isdigit() else attempt
            _collect_checks(child, data, label, checks)
            continue
        if child.tag in ("for", "if", "branch", "try", "while", "group"):
            _collect_checks(child, data, attempt, checks)
            continue
        if child.tag != "kw":
            continue
        name = child.get("name", "")
        args = {index: _resolve(arg.text, data) for index, arg in enumerate(child.findall("./arg"))}
        descriptions = _check_descriptions(data, args)
        has_children = child.find("./kw") is not None
        is_check = name in descriptions or (
            name.startswith(("Validate", "Should", "Verify")) and not has_children
        )
        status = _status(child)
        if is_check and not (name == "Fail" and status == "NOT RUN"):
            description = descriptions.get(name) or _humanize(name)
            if attempt:
                description = f"{attempt}: {description}"
            if status == "FAIL":
                detail = _failure_message(child)
            elif status == "NOT RUN":
                detail = "Not run because an earlier check failed."
            else:
                detail = _pass_detail(child)
            checks.append((description, status, detail))
        elif has_children and name not in DO_NOT_DESCEND:
            _collect_checks(child, data, attempt, checks)
    return checks


def _collect_tests(root):
    tests = []

    def visit(suite):
        for test in suite.findall("./test"):
            tests.append((suite.get("name", ""), test))
        for child in suite.findall("./suite"):
            visit(child)

    for suite in root.findall("./suite"):
        visit(suite)
    return tests


def generate_report(output_xml, destination):
    root = ET.parse(output_xml).getroot()
    tests = _collect_tests(root)

    document = Document()
    section = document.sections[0]
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(0.6))
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)

    title = document.add_heading("API Test Execution Report", 0)
    title.alignment = 1
    subtitle = document.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y, %H:%M')}")
    subtitle.alignment = 1

    # 1. Summary
    total = len(tests)
    passed = sum(_status(test) == "PASS" for _, test in tests)
    failed = sum(_status(test) == "FAIL" for _, test in tests)
    root_status = root.find("./suite/status")
    overall = "All tests passed" if failed == 0 else f"{failed} of {total} test(s) failed"
    document.add_heading("1. Execution Summary", level=1)
    summary = _table(document, ["Item", "Value"])
    for label, value in (
        ("Total tests", str(total)),
        ("Passed", (str(passed), GREEN)),
        ("Failed", (str(failed), RED if failed else GREEN)),
        ("Pass rate", f"{(passed / total * 100 if total else 0):.1f}%"),
        ("Total duration", format_duration(root_status.get("elapsed") if root_status is not None else 0)),
        ("Overall result", (overall, RED if failed else GREEN)),
    ):
        _add_row(summary, (label, value))

    # 2. Overview
    document.add_heading("2. Test Case Overview", level=1)
    overview = _table(document, ["#", "Test case", "Suite", "Status", "Duration"])
    for index, (suite_name, test) in enumerate(tests, start=1):
        status = _status(test)
        _add_row(
            overview,
            (str(index), test.get("name", ""), suite_name, _status_label(status),
             format_duration(test.find("./status").get("elapsed"))),
        )

    # 3. Per-test details
    document.add_heading("3. Test Case Details", level=1)
    for index, (suite_name, test) in enumerate(tests, start=1):
        status = _status(test)
        heading = document.add_heading(level=2)
        heading.add_run(f"3.{index}  {test.get('name', '')}  ")
        label, color = _status_label(status)
        badge = heading.add_run(f"[{label}]")
        badge.font.color.rgb = color

        info = document.add_paragraph()
        info.add_run("Suite: ").bold = True
        info.add_run(suite_name)
        info.add_run("    Duration: ").bold = True
        info.add_run(format_duration(test.find("./status").get("elapsed")))

        data = _test_data(test)
        if data:
            document.add_paragraph().add_run("Test data").bold = True
            data_table = _table(document, ["Field", "Value"])
            for key, value in data.items():
                _add_row(data_table, (DATA_LABELS.get(key, _humanize(key)), value))

        checks = _collect_checks(test, data)
        document.add_paragraph().add_run("Checks performed").bold = True
        check_table = _table(document, ["#", "Check", "Result", "Details"])
        for number, (description, check_status, detail) in enumerate(checks, start=1):
            _add_row(
                check_table,
                (str(number), description, _status_label(check_status), detail or "-"),
            )
        if not checks:
            _add_row(check_table, ("-", "No checks recorded", "-", "-"))

        if status == "FAIL":
            reason = document.add_paragraph()
            reason.add_run("Failure reason: ").bold = True
            reason_text = reason.add_run(_failure_message(test) or "-")
            reason_text.font.color.rgb = RED

    _clear_cell_widths(document)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: generate_docx_report.py <output.xml> <report.docx>")
    generate_report(Path(sys.argv[1]), Path(sys.argv[2]))
