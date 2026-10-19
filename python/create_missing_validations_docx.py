from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    """Set background color of a cell using hex string e.g. '003366'."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/>'
        f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"{title}: ")
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    run_title.font.size = Pt(10.5)
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_docx(output_path):
    doc = Document()
    
    # Page Setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("Test Review & Missing Validation Analysis")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    sub_run = sub_p.add_run("Focus File: channel_based_product_filteration.robot")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(0x46, 0x82, 0xB4)
    
    # Metadata Box
    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.style = 'Table Grid'
    meta_data = [
        [("Target Test Suite:", True), ("tests/Regression Suite/channel_based_product_filteration.robot", False)],
        [("Date of Assessment:", True), (f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", False)],
    ]
    for row_idx, row_content in enumerate(meta_data):
        for col_idx, (text, is_label) in enumerate(row_content):
            cell = meta_tbl.cell(row_idx, col_idx)
            set_cell_background(cell, "F9FAFB" if is_label else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.bold = is_label
            run.font.size = Pt(10)
            if is_label:
                run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 1. Executive Summary
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h1.runs[0].font.size = Pt(16)
    
    p = doc.add_paragraph(
        "This document presents a comprehensive review of the channel_based_product_filteration.robot test suite "
        "and its underlying validation keywords (defined in python/api_custom_validations.py and resources/keywords/api_keywords.robot). "
        "While the current implementation effectively covers basic happy-path assertions (HTTP status 200, response SLA, and nested bundled product channel matching), "
        "critical functional, schema, protocol, and negative test validations are currently missing."
    )
    p.paragraph_format.space_after = Pt(8)
    
    add_callout(
        doc,
        "A total of 12 distinct missing validations were identified across 4 key categories: Top-Level Offering Assertions, "
        "HTTP Header & Protocol Checks, TMF620 Data & Schema Integrity, and Negative/Security Test Scenarios.",
        title="KEY ASSESSMENT FINDING"
    )

    # -------------------------------------------------------------
    # 2. Existing Test Coverage Overview
    # -------------------------------------------------------------
    h1 = doc.add_heading("2. Existing Test Suite Overview", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h1.runs[0].font.size = Pt(16)
    
    p = doc.add_paragraph(
        "The channel_based_product_filteration.robot suite executes a data-driven test loop over product_offering.csv. "
        "For each record, it invokes GET /dcm/product-catalogue/v1/product-offering with parameters: "
        "relatedParty.name, depth, code, category.name, and channel.name."
    )
    p.paragraph_format.space_after = Pt(8)

    doc.add_heading("Current Executed Validations", level=2)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = table.rows[0].cells
    hdr_titles = ["Validation Name", "Source Keyword / Python Function", "Validation Scope & Behavior"]
    for idx, title in enumerate(hdr_titles):
        hdr_cells[idx].text = title
        set_cell_background(hdr_cells[idx], "003366")
        set_cell_margins(hdr_cells[idx], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[idx].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10)

    existing_rows = [
        ("Validate Response Status", "validate_response_status(200)", "Asserts HTTP status code equals 200 OK."),
        ("Validate Response Not Empty", "validate_response_not_empty()", "Asserts response body is non-empty."),
        ("Validate Response Time", "validate_response_time(5000)", "Asserts response latency is <= 5000 milliseconds."),
        ("Validate Product Offering Response", "validate_product_offering_response()", "Asserts code, relatedParty, and category match expected CSV row values."),
        ("Validate Bundled Product Offering Channels", "validate_bundled_product_offering_channels()", "Recursively validates that all items in bundledProductOffering list contain requested channel.name.")
    ]

    for v_name, v_src, v_desc in existing_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = v_name
        row_cells[1].text = v_src
        row_cells[2].text = v_desc
        for c in row_cells:
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            c.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 3. Detailed Missing Validations & Functional Gaps
    # -------------------------------------------------------------
    h1 = doc.add_heading("3. Identified Missing Validations & Functional Gaps", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h1.runs[0].font.size = Pt(16)
    
    categories = [
        ("A. Top-Level Offering & Pricing Channel Validations", [
            ("1. Parent Product Offering Channel Validation", 
             "HIGH",
             "Current Gap: Validate Bundled Product Offering Channels checks channels inside bundledProductOffering arrays, but does NOT verify that the root/parent product offering's channel array includes requested channel.name.\n"
             "Risk: A parent product offering returned for a restricted channel (e.g. 'WEB') might have an invalid or missing top-level channel filter in the backend API response."),
            ("2. Direct Execution of Price Channel Validation", 
             "MEDIUM",
             "Current Gap: Helper functions exist in Python (_validate_price_channels), but product offering price channel assertions are not explicitly reported or structured as distinct sub-validations in the Robot execution log.\n"
             "Risk: Invalid pricing channel mapping on the root offering passes undetected if prices list is null or improperly structured.")
        ]),
        ("B. HTTP Protocol & Header Integrity Validations", [
            ("3. Content-Type Header Verification", 
             "HIGH",
             "Current Gap: The test does not verify that the HTTP response Content-Type header matches 'application/json' or 'application/json;charset=UTF-8'.\n"
             "Risk: API returning HTML error pages or text/plain payloads with status 200 OK could cause unexpected parsing failures."),
            ("4. Cache-Control & Pragma Header Verification", 
             "LOW",
             "Current Gap: No verification of Cache-Control header (e.g. 'no-store', 'no-cache') or custom version headers.\n"
             "Risk: Product catalog pricing/offering data might be accidentally cached by downstream API gateways or proxies.")
        ]),
        ("C. TMF620 Schema & Data Integrity Validations", [
            ("5. Lifecycle Status Validation (lifecycleStatus)", 
             "HIGH",
             "Current Gap: The test does not check if lifecycleStatus is in an active state (e.g., 'Active', 'Launched').\n"
             "Risk: Obsolete, Retired, or Draft product offerings may leak into customer-facing sales channels."),
            ("6. Price Object & Currency Structure Integrity", 
             "MEDIUM",
             "Current Gap: Does not validate productOfferingPrice sub-fields such as priceType ('recurring', 'one-time'), price.value (> 0), price.unit, taxIncluded, or recurringChargePeriod.\n"
             "Risk: Zero/negative price values or missing currency codes will pass without failing the test."),
            ("7. Validity Period (validFor) Date Verification", 
             "MEDIUM",
             "Current Gap: validFor.startDateTime and validFor.endDateTime are not validated against current system time.\n"
             "Risk: Expired or future-dated product offerings may be incorrectly returned as active."),
            ("8. Mandatory TMF620 Attribute Presence", 
             "MEDIUM",
             "Current Gap: Missing explicit schema checks for standard TMF620 fields: id, href, name, isBundle, @type, @baseType.\n"
             "Risk: Structurally incomplete API payloads missing mandatory TMF fields pass as successful.")
        ]),
        ("D. Negative Testing, Boundary Conditions & Authorization", [
            ("9. Non-Existent or Invalid Channel Filtering (Negative Path)", 
             "HIGH",
             "Current Gap: No test cases pass an invalid channel name (e.g. 'INVALID_CHANNEL_999').\n"
             "Expected Behavior: API should return HTTP 404 Not Found or HTTP 200 with an empty list []. Currently untested."),
            ("10. Invalid Product Code Query Handling", 
             "MEDIUM",
             "Current Gap: No test cases pass an invalid/non-existent product code.\n"
             "Expected Behavior: Should validate correct error structure or HTTP 404 response."),
            ("11. Unauthorized Access & Expired Token Testing", 
             "HIGH",
             "Current Gap: No test case sends request without Bearer token or with an invalid token.\n"
             "Expected Behavior: Should validate HTTP 401 Unauthorized."),
            ("12. Query Parameter Field Selection (fields=...) & Pagination", 
             "LOW",
             "Current Gap: Query parameters like fields (e.g. fields=id,name,code) and limit/offset are defined in keyword parameters but never asserted in tests.\n"
             "Risk: Unrestricted field payloads wasting bandwidth when restricted fields are requested.")
        ])
    ]

    for cat_title, items in categories:
        h2 = doc.add_heading(cat_title, level=2)
        h2.runs[0].font.color.rgb = RGBColor(0x46, 0x82, 0xB4)
        
        for item_title, priority, details in items:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            
            run_prio = p.add_run(f"[{priority}] ")
            run_prio.bold = True
            if priority == "HIGH":
                run_prio.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            elif priority == "MEDIUM":
                run_prio.font.color.rgb = RGBColor(0xD9, 0x82, 0x2B)
            else:
                run_prio.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
                
            run_t = p.add_run(item_title)
            run_t.bold = True
            run_t.font.size = Pt(11)
            
            p_desc = doc.add_paragraph(details)
            p_desc.paragraph_format.left_indent = Inches(0.2)
            p_desc.paragraph_format.space_after = Pt(6)
            p_desc.runs[0].font.size = Pt(10)

    # -------------------------------------------------------------
    # 4. Matrix of Validation Coverage (Existing vs Missing)
    # -------------------------------------------------------------
    h1 = doc.add_heading("4. Validation Matrix: Existing vs. Recommended", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h1.runs[0].font.size = Pt(16)
    
    matrix_tbl = doc.add_table(rows=1, cols=4)
    matrix_tbl.style = 'Table Grid'
    matrix_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    m_headers = ["Validation Target", "Current Status", "Priority", "Recommended Assertion"]
    for idx, title in enumerate(m_headers):
        cell = matrix_tbl.rows[0].cells[idx]
        cell.text = title
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10)

    matrix_data = [
        ("HTTP Status 200", "COVERED", "LOW", "validate_response_status(200)"),
        ("Response Body Non-Empty", "COVERED", "LOW", "validate_response_not_empty()"),
        ("Response SLA (<5000ms)", "COVERED", "LOW", "validate_response_time(5000)"),
        ("Bundled Offering Channel", "COVERED", "MEDIUM", "Validate Bundled Product Offering Channels"),
        ("Root Product Channel", "MISSING", "HIGH", "Assert requested channel in parent offering.channel[]"),
        ("Content-Type Header", "MISSING", "HIGH", "Assert Content-Type == application/json"),
        ("lifecycleStatus", "MISSING", "HIGH", "Assert lifecycleStatus in ['Active', 'Launched']"),
        ("Price Structure & Value", "MISSING", "MEDIUM", "Assert price.value > 0 and priceType exists"),
        ("validFor Date Range", "MISSING", "MEDIUM", "Assert startDateTime <= current <= endDateTime"),
        ("Mandatory TMF Fields", "MISSING", "MEDIUM", "Assert id, href, name, @type exist"),
        ("Invalid Channel (Negative)", "MISSING", "HIGH", "Assert 404 or empty array [] for invalid channel"),
        ("Unauthorized Access", "MISSING", "HIGH", "Assert HTTP 401 when token is missing/invalid"),
    ]

    for target, status, priority, rec in matrix_data:
        row_cells = matrix_tbl.add_row().cells
        row_cells[0].text = target
        row_cells[1].text = status
        row_cells[2].text = priority
        row_cells[3].text = rec
        
        # Color coding status
        p_stat = row_cells[1].paragraphs[0]
        if status == "COVERED":
            p_stat.runs[0].font.color.rgb = RGBColor(0x00, 0x80, 0x00)
            p_stat.runs[0].bold = True
        else:
            p_stat.runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            p_stat.runs[0].bold = True
            
        for c in row_cells:
            set_cell_margins(c, top=80, bottom=80, left=80, right=80)
            c.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 5. Recommended Code Enhancements
    # -------------------------------------------------------------
    h1 = doc.add_heading("5. Recommended Implementation Snippets", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    h1.runs[0].font.size = Pt(16)
    
    doc.add_heading("A. Enhanced Python Validation Function (api_custom_validations.py)", level=2)
    
    py_code = (
        "def validate_complete_product_offering(\n"
        "    response, requested_channel, code=None, related_party=None, category_name=None\n"
        "):\n"
        "    # 1. Validate HTTP Status & Response Headers\n"
        "    validate_response_status(response, 200)\n"
        "    content_type = response.headers.get('Content-Type', '')\n"
        "    if 'application/json' not in content_type:\n"
        "        raise AssertionError(f'Expected Content-Type application/json, got: {content_type}')\n\n"
        "    # 2. Parse Payload and Extract Offerings\n"
        "    payload = _get_json_payload(response)\n"
        "    offerings = _get_product_offerings(payload)\n\n"
        "    for offering in offerings:\n"
        "        # 3. Validate Mandatory TMF620 Attributes\n"
        "        for field in ['id', 'name', 'code', 'lifecycleStatus']:\n"
        "            if field not in offering:\n"
        "                raise AssertionError(f'Offering missing mandatory field: {field}')\n\n"
        "        # 4. Validate Lifecycle Status\n"
        "        status = offering.get('lifecycleStatus')\n"
        "        if status not in ['Active', 'Launched']:\n"
        "            raise AssertionError(f'Offering {offering.get(\"code\")} has invalid lifecycleStatus: {status}')\n\n"
        "        # 5. Validate Root Offering Channel Presence\n"
        "        channels, issue = _get_available_channels(offering)\n"
        "        normalized = {c.casefold() for c in channels}\n"
        "        if requested_channel.strip().casefold() not in normalized:\n"
        "            raise AssertionError(\n"
        "                f'Root offering {offering.get(\"code\")} missing requested channel \"{requested_channel}\". '\n"
        "                f'Available channels: {channels}'\n"
        "            )\n"
    )
    
    code_tbl = doc.add_table(rows=1, cols=1)
    code_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = code_tbl.cell(0, 0)
    set_cell_background(c_cell, "F5F5F5")
    set_cell_margins(c_cell, top=100, bottom=100, left=150, right=150)
    p_code = c_cell.paragraphs[0]
    p_code.paragraph_format.space_after = Pt(0)
    r_code = p_code.add_run(py_code)
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(9)
    r_code.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.add_heading("B. Updated Robot Framework Test Case (channel_based_product_filteration.robot)", level=2)
    
    robot_code = (
        "*** Test Cases ***\n"
        "Get Product Offering From CSV\n"
        "    ${rows}=    Read Csv    ${CSV_FILE}\n"
        "    FOR    ${row}    IN    @{rows}\n"
        "        ${response}=    Get Product Offering    ...\n"
        "        \n"
        "        # Comprehensive Assertions\n"
        "        Validate Response Status           ${response}    200\n"
        "        Validate Response Header           ${response}    Content-Type    application/json\n"
        "        Validate Response Time             ${response}    5000\n"
        "        Validate Complete Product Offering ${response}    requested_channel=${channel_name} ...\n"
        "        Validate Bundled Product Offering Channels    ${response}    ${channel_name}\n"
        "    END\n"
    )
    
    code_tbl2 = doc.add_table(rows=1, cols=1)
    code_tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell2 = code_tbl2.cell(0, 0)
    set_cell_background(c_cell2, "F5F5F5")
    set_cell_margins(c_cell2, top=100, bottom=100, left=150, right=150)
    p_code2 = c_cell2.paragraphs[0]
    p_code2.paragraph_format.space_after = Pt(0)
    r_code2 = p_code2.add_run(robot_code)
    r_code2.font.name = 'Consolas'
    r_code2.font.size = Pt(9)
    r_code2.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)

    doc.save(output_path)
    print(f"Successfully generated DOCX report at: {output_path}")

if __name__ == '__main__':
    out_dir = Path(__file__).resolve().parent.parent / "reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / "Channel_Based_Product_Filteration_Missing_Validations_Report.docx"
    build_docx(out_file)
